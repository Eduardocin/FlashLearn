from .models import UserFlashcard, ReviewLog
from ai.api import generate_flashcards
import fitz
import chardet
import docx
from fpdf import FPDF
from io import BytesIO
from django.utils import timezone
from datetime import timedelta
from collections import defaultdict

# ─── Intervalos SM-2 simplificado ────────────────────────────────────────
# Acertos consecutivos: 0→1d, 1→3d, 2→7d, 3→14d, 4+→30d
SR_INTERVALS = [1, 3, 7, 14, 30]


class SpacedRepetitionService:
    """Calcula status de repetição espaçada para flashcards."""

    @staticmethod
    def _next_interval(consecutive_correct: int) -> int:
        """Retorna o intervalo em dias baseado nos acertos consecutivos."""
        idx = min(consecutive_correct, len(SR_INTERVALS) - 1)
        return SR_INTERVALS[idx]

    @staticmethod
    def card_sr_info(flashcard: UserFlashcard) -> dict:
        """
        Retorna dict com info SR de um único flashcard:
          - last_reviewed: datetime ou None
          - due_date: datetime da próxima revisão
          - status: 'never' | 'overdue' | 'due_today' | 'upcoming'
          - days_until_due: int (negativo = atrasado)
          - consecutive_correct: int
          - total_reviews: int
          - accuracy: float 0-100
        """
        logs = list(
            ReviewLog.objects.filter(flashcard=flashcard)
            .order_by('-reviewed_at')
            .values('is_correct', 'reviewed_at')
        )

        now = timezone.now()

        if not logs:
            return {
                'last_reviewed': None,
                'due_date': now,
                'status': 'never',
                'days_until_due': 0,
                'consecutive_correct': 0,
                'total_reviews': 0,
                'accuracy': 0.0,
            }

        # Contar acertos consecutivos a partir do mais recente
        consecutive_correct = 0
        for log in logs:
            if log['is_correct']:
                consecutive_correct += 1
            else:
                break  # interrompe na primeira errada

        last_reviewed = logs[0]['reviewed_at']

        # Se o último foi errado → intervalo = 1 dia
        if not logs[0]['is_correct']:
            consecutive_correct = 0

        interval_days = SpacedRepetitionService._next_interval(consecutive_correct)
        due_date = last_reviewed + timedelta(days=interval_days)
        days_until_due = (due_date.date() - now.date()).days

        if days_until_due < 0:
            status = 'overdue'
        elif days_until_due == 0:
            status = 'due_today'
        else:
            status = 'upcoming'

        total_reviews = len(logs)
        correct_count = sum(1 for l in logs if l['is_correct'])
        accuracy = round(correct_count / total_reviews * 100, 1) if total_reviews else 0.0

        return {
            'last_reviewed': last_reviewed,
            'due_date': due_date,
            'status': status,
            'days_until_due': days_until_due,
            'consecutive_correct': consecutive_correct,
            'total_reviews': total_reviews,
            'accuracy': accuracy,
        }

    @staticmethod
    def session_sr_summary(flashcards: list) -> dict:
        """
        Recebe lista de UserFlashcard de uma sessão e retorna resumo SR:
          - overdue: int
          - due_today: int
          - upcoming: int
          - never: int
          - next_due: date (a mais próxima)
          - last_reviewed: datetime (a mais recente)
          - overall_accuracy: float
        """
        overdue = due_today = upcoming = never = 0
        next_due = None
        last_reviewed = None
        accuracies = []

        for card in flashcards:
            info = SpacedRepetitionService.card_sr_info(card)
            s = info['status']
            if s == 'overdue':
                overdue += 1
            elif s == 'due_today':
                due_today += 1
            elif s == 'upcoming':
                upcoming += 1
            else:
                never += 1

            if info['last_reviewed']:
                if last_reviewed is None or info['last_reviewed'] > last_reviewed:
                    last_reviewed = info['last_reviewed']

            if s != 'never':
                due = info['due_date']
                if next_due is None or due < next_due:
                    next_due = due

            if info['total_reviews'] > 0:
                accuracies.append(info['accuracy'])

        overall_accuracy = round(sum(accuracies) / len(accuracies), 1) if accuracies else None

        return {
            'overdue': overdue,
            'due_today': due_today,
            'upcoming': upcoming,
            'never': never,
            'next_due': next_due,
            'last_reviewed': last_reviewed,
            'overall_accuracy': overall_accuracy,
        }

class FlashcardService:
    @staticmethod
    def extract_text_from_file(file):
        """
        Extract text from a file based on its extension.
        The supported extensions are: .pdf, .docx and .txt.
        """
        if file.name.lower().endswith(".pdf"):
            return FlashcardService._extract_text_from_pdf(file)
        elif file.name.lower().endswith(".docx"):
            return FlashcardService._extract_text_from_docx(file)
        else:
            return FlashcardService._extract_text_from_txt(file)
        
        
    @staticmethod
    def _extract_text_from_pdf(file):
        text = ""
        with fitz.open(stream=file.read(), filetype="pdf") as doc:
            for page in doc:
                text += page.get_text()
        return text[:3000]
    
    @staticmethod
    def _extract_text_from_docx(file):
        doc = docx.Document(file)
        return '\n'.join([p.text for p in doc.paragraphs])[:3000]
    
    @staticmethod
    def _extract_text_from_txt(file):
        raw_content = file.read()
        encoding = chardet.detect(raw_content).get('encoding')
        file.seek(0)
        return raw_content.decode(encoding)[:3000]
    
    @staticmethod
    def create_flashcards_from_file(user, file, collection=None):
        text = FlashcardService.extract_text_from_file(file)
        api_flashcards = generate_flashcards(text)
        
        db_flashcards = []
        for data in api_flashcards:
            card = UserFlashcard.objects.create(
                user=user,
                title=data['title'],
                content=data['content'],
                collection=collection,
            )
            db_flashcards.append({
                'id': card.id,
                'title': card.title,
                'content': card.content
            })
        
        return db_flashcards


    @staticmethod
    def update_flashcards_from_data(user, titles, contents, flashcard_ids):
        updated_flashcards = []
        
        for i in range(len(titles)):
            if i >= len(contents): 
                continue
                
            title = titles[i]
            content = contents[i]
            
            card_id = None
            if i < len(flashcard_ids) and flashcard_ids[i] and flashcard_ids[i].strip():
                try:
                    card_id = int(flashcard_ids[i])
                except (ValueError, TypeError):
                    card_id = None
            
            if card_id:
                card, _ = UserFlashcard.objects.update_or_create(
                    id=card_id, 
                    user=user,
                    defaults={'title': title, 
                                'content': content}
                )
            else:
                card = UserFlashcard.objects.create(
                    user=user, 
                    title=title, 
                    content=content
                )
            
            updated_flashcards.append({
                'id': card.id,
                'title': card.title, 
                'content': card.content
            })
        
        return updated_flashcards

class PDFService:
    @staticmethod
    def generate_flashcards_pdf(flashcards):
        if not flashcards:
            return None
        
        buffer_pdf = BytesIO()
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        
        pdf.set_font("Arial", 'B', 16)
        pdf.cell(0, 10, txt="Flashcards Gerados", ln=1, align='C')
        pdf.ln(10)
        
        
        pdf.set_font("Arial", size=12)
        for idx, card in enumerate(flashcards, 1):
            # Título
            pdf.set_font("Arial", 'B', 12)
            titulo = f"{idx}. {card.title}"
            safe_titulo = titulo.encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 10, txt=safe_titulo)
            
            # Conteúdo
            pdf.set_font("Arial", '', 12)
            conteudo = f"{card.content}"
            safe_conteudo = conteudo.encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 10, txt=safe_conteudo)
            
            pdf.ln(5)
        
        pdf_content = pdf.output(dest='S').encode('latin-1')
        buffer_pdf.write(pdf_content)
        buffer_pdf.seek(0)
        return buffer_pdf
    